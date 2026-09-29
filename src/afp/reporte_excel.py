"""Reporte maestro en Excel: resultados + simulador de pensión 100% en fórmulas."""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import config, mortalidad

F = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
AMARILLO = PatternFill("solid", fgColor="FFFF00")
AZUL, VERDE = "0000FF", "008000"


def _f(**k):
    return Font(name=F, size=10, **k)


def _encabezado(ws, fila, valores, anchos=None):
    for i, v in enumerate(valores, 1):
        c = ws.cell(fila, i, v)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF"); c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(anchos or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _tabla(wb, nombre, df, formatos=None, titulo=None, notas=(), anchos=None):
    ws = wb.create_sheet(nombre)
    fila = 1
    if titulo:
        ws.cell(1, 1, titulo).font = Font(name=F, size=13, bold=True); fila = 3
    _encabezado(ws, fila, list(df.columns), anchos or [max(12, min(34, len(str(c)) + 4)) for c in df.columns])
    for r in df.itertuples(index=False):
        fila += 1
        for j, v in enumerate(r, 1):
            c = ws.cell(fila, j, v.item() if hasattr(v, "item") else v)
            c.font = _f(color=AZUL) if isinstance(v, (int, float)) else _f()
            fmt = (formatos or {}).get(df.columns[j - 1])
            if fmt:
                c.number_format = fmt
    for i, n in enumerate(notas, fila + 2):
        ws.cell(i, 1, n).font = _f(italic=True)
    ws.freeze_panes = ws.cell(4 if titulo else 2, 1)
    return ws


def construir(ruta, r: dict):
    wb = Workbook(); wb.calculation.fullCalcOnLoad = True

    # --- Resumen ------------------------------------------------------------
    ws = wb.active; ws.title = "Resumen"
    ws["A1"] = "De multifondos a fondos generacionales: riesgo, regímenes y pensiones en Chile"
    ws["A1"].font = Font(name=F, size=14, bold=True)
    ws["A2"] = f"Datos: {r['rango_datos']}. Valores en azul provienen del pipeline en Python (run_pipeline.py)."
    ws["A2"].font = _f(italic=True)
    _encabezado(ws, 4, ["#", "Hallazgo", "Evidencia"], [5, 70, 70])
    for i, (h, e) in enumerate(r["hallazgos"], 5):
        ws.cell(i, 1, i - 4).font = _f(); ws.cell(i, 2, h).font = _f(bold=True); ws.cell(i, 3, e).font = _f()
        ws.cell(i, 2).alignment = ws.cell(i, 3).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 48

    # --- Parámetros (inputs del simulador) ---------------------------------
    P = wb.create_sheet("Parametros")
    _encabezado(P, 1, ["Parámetro", "Valor", "Descripción / fuente"], [30, 16, 90])
    params = [("Sexo (H/M)", "H", "H usa CB-H-2020; M usa RV-M-2020 (CMF/SP, NCG 495, 2023)"),
              ("Edad al pensionarse", 65, "Edad legal: 65 hombres, 60 mujeres"),
              ("Saldo acumulado (UF)", 3000, "Saldo de la cuenta individual al pensionarse"),
              ("Año de cálculo", 2026, "Año desde el que se proyecta la mortalidad"),
              ("TITRP anual", config.TITRP, f"Tasa técnica de retiro programado, vigente desde {config.TITRP_VIGENCIA} (Circular SP N.º 2417)"),
              ("Valor UF (CLP)", r["uf_ultima"], f"UF al {r['uf_fecha']} (Banco Central)")]
    for i, (a, b, c) in enumerate(params, 2):
        P.cell(i, 1, a).font = _f(); x = P.cell(i, 2, b); x.font = _f(color=AZUL); x.fill = AMARILLO
        P.cell(i, 3, c).font = _f()
    P["B6"].number_format = "0.00%"; P["B7"].number_format = "#,##0.00"; P["B4"].number_format = "#,##0"
    dv = DataValidation(type="list", formula1='"H,M"'); P.add_data_validation(dv); dv.add("B2")
    P["A9"] = "Celda amarilla con texto azul = editable. Negro = fórmula. Verde = vínculo a otra hoja."
    P["A9"].font = _f(italic=True)
    P["A11"] = "Supuestos de la simulación Monte Carlo (config.py)"; P["A11"].font = _f(bold=True)
    for i, (k, v) in enumerate(config.SIM.items(), 12):
        P.cell(i, 1, k).font = _f(); P.cell(i, 2, v).font = _f(color=AZUL)

    # --- Mortalidad proyectada (datos para el simulador) --------------------
    anios = list(range(2026, 2101))
    for sexo, nombre, tabla in (("H", "Mortalidad_H", "CB-H-2020"), ("M", "Mortalidad_M", "RV-M-2020")):
        edades, _, M = mortalidad.matriz_qx(sexo, anios)
        w = wb.create_sheet(nombre)
        w.cell(1, 1, f"qx proyectado = qx_2020 × Π(1 − AA x,t), tabla {tabla}. Filas: edad; columnas: año. "
                     "Desde 2036 el factor AA se mantiene constante.").font = _f(italic=True)
        w.cell(2, 1, "Edad").font = _f(bold=True)
        for j, a in enumerate(anios, 2):
            w.cell(2, j, a).font = _f(bold=True)
        for i, e in enumerate(edades, 3):
            w.cell(i, 1, e)
            for j in range(len(anios)):
                w.cell(i, j + 2, round(float(M[i - 3, j]), 8)).number_format = "0.000000"
        w.freeze_panes = "B3"
    ultima_fila_h = 2 + len(mortalidad.matriz_qx("H", [2026])[0])
    ultima_fila_m = 2 + len(mortalidad.matriz_qx("M", [2026])[0])
    col_fin = get_column_letter(1 + len(anios))

    # --- Simulador ----------------------------------------------------------
    S = wb.create_sheet("Simulador_Pension", 1)
    S["A1"] = "Simulador de pensión por Retiro Programado (primer año)"; S["A1"].font = Font(name=F, size=14, bold=True)
    S["A2"] = "Cambia los valores amarillos en 'Parametros'. Simplificación: sin beneficiarios de sobrevivencia."
    S["A2"].font = _f(italic=True)
    salidas = [("Sexo", "=Parametros!B2", None), ("Edad al pensionarse", "=Parametros!B3", "0"),
               ("Saldo (UF)", "=Parametros!B4", "#,##0"), ("TITRP", "=Parametros!B6", "0.00%"),
               ("Capital Necesario Unitario (CNU)", "=SUM(H12:H731)", "#,##0.00"),
               ("Pensión mensual (UF)", "=C6/C8", "#,##0.00"),
               ("Pensión mensual (CLP)", "=C9*Parametros!B7", "$#,##0")]
    for i, (a, b, fmt) in enumerate(salidas, 4):
        S.cell(i, 2, a).font = _f(bold=i >= 8)
        c = S.cell(i, 3, b); c.font = _f(color=VERDE if "Parametros" in b else "000000", bold=i >= 8)
        if fmt:
            c.number_format = fmt
    _encabezado(S, 11, ["Mes k", "Edad", "Año", "qx anual", "q mensual", "Prob. vivo al pago",
                        "Factor descuento", "VP pago"], [10, 34, 12, 12, 12, 18, 16, 14])
    rng_h = f"Mortalidad_H!$B$3:${col_fin}${ultima_fila_h}"; rng_m = f"Mortalidad_M!$B$3:${col_fin}${ultima_fila_m}"
    for k in range(720):
        f = 12 + k
        S.cell(f, 1, k).font = _f(color=AZUL)
        S.cell(f, 2, f"=$C$5+INT(A{f}/12)")
        S.cell(f, 3, f"=Parametros!$B$5+INT(A{f}/12)")
        S.cell(f, 4, f'=IFERROR(IF($C$4="H",INDEX({rng_h},MATCH(B{f},Mortalidad_H!$A$3:$A${ultima_fila_h},0),'
                     f'MATCH(MIN(C{f},2100),Mortalidad_H!$B$2:${col_fin}$2,0)),INDEX({rng_m},'
                     f'MATCH(B{f},Mortalidad_M!$A$3:$A${ultima_fila_m},0),MATCH(MIN(C{f},2100),Mortalidad_M!$B$2:${col_fin}$2,0))),1)')
        S.cell(f, 5, f"=1-(1-D{f})^(1/12)")
        S.cell(f, 6, 1 if k == 0 else f"=F{f - 1}*(1-E{f - 1})")
        S.cell(f, 7, f"=1/(1+$C$7)^(A{f}/12)")
        S.cell(f, 8, f"=F{f}*G{f}")
        for c, fmt in ((4, "0.000000"), (5, "0.000000"), (6, "0.0000"), (7, "0.0000"), (8, "0.0000")):
            S.cell(f, c).number_format = fmt
    S.freeze_panes = "A12"

    pct = "0.00%"
    _tabla(wb, "Calidad_datos", r["calidad"], titulo="Validaciones de la base de valores cuota (etapa 1)",
           notas=["Rebases por fusión detectados y neutralizados:"] +
                 [f"  {e.fecha:%Y-%m-%d} fondo {e.fondo} {e.afp}: +{e.retorno_bruto:.1%} ({e.diagnostico}, desaparece {e.afp_que_desaparece})"
                  for e in r["eventos"].itertuples()])
    _tabla(wb, "Rentabilidad", r["rentabilidad"],
           {c: pct for c in r["rentabilidad"].columns if c != "fondo"},
           titulo="Rentabilidad y riesgo por fondo, ago-2002 a ago-2026 (ponderado por afiliados)",
           notas=["Real = valor cuota en UF. La inflación implícita es igual en los 5 fondos: control de consistencia."])
    _tabla(wb, "Regimenes_HMM", r["regimenes"], {"retorno_anual": pct, "vol_anual": pct, "prob_estacionaria": pct,
                                                "duracion_esperada_meses": "0.0"},
           titulo="HMM gaussiano de 3 estados sobre retornos reales mensuales de los 5 fondos",
           notas=[f"Selección por BIC: {r['bic_texto']}", "Matriz de transición (fila = desde):"] +
                 ["  " + "  ".join(f"{x:.3f}" for x in fila) for fila in r["transicion"]])
    _tabla(wb, "GARCH", r["garch"], {"alpha": "0.0000", "beta": "0.0000", "persistencia": "0.0000",
                                     "vol_largo_plazo_anual": pct, "vol_actual_anual": pct, "vida_media_dias": "0.0"},
           titulo="GARCH(1,1) sobre retornos reales diarios del sistema (QMLE)",
           notas=["Persistencia ~1 en D y E: huella de un cambio de nivel de volatilidad (quiebre 2019), no de shocks transitorios."])
    fm = {c: "0.00" for c in ("p5", "p25", "mediana", "p75", "p95", "media", "cvar5")}
    fm.update(tasa_reemplazo_mediana=pct, prob_tr_meta=pct)
    for sexo in ("H", "M"):
        t = r["montecarlo"][sexo]
        _tabla(wb, f"MonteCarlo_{sexo}", t, fm,
               titulo=f"Pensión mensual simulada (UF) por estrategia · {'Hombre 25-65' if sexo == 'H' else 'Mujer 25-60'}",
               notas=[f"CNU al pensionarse: {t.attrs['cnu']:.1f} · año de retiro {t.attrs['anio_retiro']} · "
                      f"sueldo final {t.attrs['sueldo_final']:.1f} UF · traspasos (mediana, reactivo): {t.attrs['cambios_reactivo_mediana']:.0f}",
                      "cvar5 = pensión promedio en el 5% de peores escenarios."])
    _tabla(wb, "Sensibilidad", r["sensibilidad"], fm,
           titulo="Sensibilidad: rentabilidades históricas castigadas en 1 p.p. anual (hombre)")
    _tabla(wb, "Backtest_2002_2026", r["backtest"], {"saldo_final_uf": "#,##0.0", "pension_mensual_uf": "0.00", "vs_defecto": pct},
           titulo="Evaluación retrospectiva con la historia real (sep-2002 a ago-2026)",
           notas=["Un solo camino histórico: útil como evidencia, no como probabilidad."])
    _tabla(wb, "Bandas_FG", r["bandas"], {"banda": pct, "pct_meses_sobre": pct, "pct_meses_bajo": pct,
                                          "pct_meses_fuera": pct, "desv_p95_abs": pct},
           titulo="Frecuencia histórica fuera de banda (AFP vs promedio de sus pares, 36 meses móviles)",
           notas=["Bandas oficiales: 240 a 290 p.b. anuales sobre la cartera de referencia (SP, Res. Ex. 1195, sep-2026).",
                  "Transición: banda de 400 p.b. (±4,0%) entre el 01-04-2028 y el 31-03-2031.",
                  "Proxy: la referencia oficial son índices de mercado, no el promedio de pares."])
    if "robustez" in r:
        _tabla(wb, "Robustez", r["robustez"],
               {c: pct for c in r["robustez"].columns if c.startswith(("cv_", "reactivo_", "prob_"))} |
               {"mediana_defecto_uf": "0.00", "traspasos_reactivo": "0"},
               titulo=f"Robustez: las conclusiones con {len(r['robustez'])} supuestos alternativos (10.000 escenarios por caso)",
               notas=["prob_* = fracción de escenarios en que, sobre la MISMA trayectoria de mercado, una estrategia supera o pierde frente a la otra.",
                      "Densidad y nivel de sueldo no cambian el ranking (la pensión es lineal en los aportes), por eso no están en la grilla."],
               anchos=[32, 18, 16, 18, 20, 18, 16, 16, 14])
    if "esg" in r:
        _tabla(wb, "Escenarios_mercado", r["esg"].astype(object).where(r["esg"].notna(), None),
               {c: pct for c in r["esg"].columns if c.startswith(("cv_", "reactivo_", "prob_", "retorno_"))} |
               {"mediana_defecto_uf": "0.00", "mediana_cv_uf": "0.00"},
               titulo="Pensión bajo supuestos de mercado forward-looking (10.000 escenarios por caso)",
               notas=["Cada fondo se recentra en r = g·r_crecimiento + (1−g)·r_protección, con g = activos de crecimiento del fondo.",
                      "Consenso 2026: protección 2,52% real (bono en UF a 10 años, BCCh); crecimiento 4,4% real (MSCI ACWI, J.P. Morgan LTCMA 2026).",
                      "Se conservan regímenes, volatilidades, colas y correlaciones históricas: solo cambia el retorno esperado."],
               anchos=[28, 6, 14, 14, 16, 14, 16, 18, 14, 12, 14, 12])
    if "reforma" in r:
        _tabla(wb, "Pension_total", r["reforma"].reindex(columns=["sexo", "estrategia", "capa", "p5", "mediana", "p95",
                                                                 "tasa_reemplazo_mediana", "prob_tr_meta"]),
               {"p5": "0.00", "mediana": "0.00", "p95": "0.00", "tasa_reemplazo_mediana": pct, "prob_tr_meta": pct},
               titulo="Pensión total con la reforma (Ley 21.735): capas acumuladas, supuestos de consenso 2026 (UF de hoy)",
               notas=["Capas: 10% del trabajador; + cotización del empleador a la cuenta individual (0,1% → 6% según fecha); "
                      "+ Cotización con Rentabilidad Protegida (bono a 2,52% real); + PGU ($250.275 feb-2026, focalizada por pensión base).",
                      "Mujer: se pensiona a los 60 y recibe la PGU desde los 65. Supone que cumple la focalización (no es del 10% más rico).",
                      "Tasa de reemplazo = pensión / último sueldo imponible. Meta: 40%."],
               anchos=[6, 26, 28, 10, 10, 10, 14, 14])
        _tabla(wb, "Pension_por_ingreso", r["reforma_ingreso"],
               {c: "0.00" for c in ["sueldo_final_uf", "autofinanciada_10", "empleador_cci", "crp", "pgu", "total"]} |
               {"tasa_reemplazo_total": pct, "tasa_reemplazo_10": pct},
               titulo="Progresividad: pensión total mediana por nivel de sueldo (fondo por defecto, consenso 2026)")
    if "optimo" in r:
        _tabla(wb, "Glidepath_optimo", r["optimo"],
               {"g0": pct, "g1": pct, "a1": "0", "ec_oficial": "0.00", "ec_optimo": "0.00", "ec_defecto": "0.00",
                "optimo_vs_oficial": pct, "defecto_vs_oficial": pct, "gamma": "0.0"},
               titulo="¿Es óptimo el glidepath oficial? Mejor glidepath de la familia por aversión al riesgo (equivalente cierto, UF)",
               notas=["Familia: meseta g0 hasta la edad a1 y baja lineal hasta g1 al retiro; réplica con multifondos.",
                      "Criterio: utilidad CRRA de la pensión total (con_reforma = cotización del empleador, rentabilidad protegida y PGU).",
                      "Grilla evaluada en escenarios de entrenamiento; resultados reportados en escenarios de prueba independientes."])
    if "incertidumbre" in r:
        t = r["incertidumbre"].reindex(columns=["bloque", "regimen", "fondo", "metrica", "p05", "mediana", "p95"])
        t = t.replace([np.inf, -np.inf], np.nan).astype(object).where(t.notna(), None)
        _tabla(wb, "Incertidumbre", t, {"p05": "0.0000", "mediana": "0.0000", "p95": "0.0000"},
               titulo="Incertidumbre de parámetros: intervalos de confianza de 90% por bootstrap (percentiles 5–95)",
               notas=["Pensiones: bootstrap estacionario de la historia (bloques de 12 meses en promedio), reestimando el HMM "
                      "y repitiendo el Monte Carlo (hombre, 2.000 escenarios por historia).",
                      "HMM: bootstrap paramétrico con estados alineados a los originales (asignación húngara).",
                      "GARCH: simulación histórica filtrada (residuos estandarizados remuestreados). Persistencia 0,9999 = "
                      "tope de la restricción (volatilidad integrada, vida media no acotada)."],
               anchos=[34, 9, 7, 30, 12, 12, 12])
    fu = wb.create_sheet("Fuentes")
    for i, s in enumerate(r["fuentes"], 1):
        fu.cell(i, 1, s).font = _f()
    fu.column_dimensions["A"].width = 140
    wb.save(ruta)
    return ruta
